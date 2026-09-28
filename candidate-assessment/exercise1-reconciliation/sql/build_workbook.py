"""Write Reconciliation_Workbook.xlsx (called by 07_export_exceptions.py).

Sheets: Summary (bridge + category summary, all formulas), Exceptions, Internal deposits,
Gateway settlement (the raw files), Assumptions. Every total on Summary is a SUMIFS/COUNTIFS on
the other sheets, so the arithmetic can be audited and re-run if a row changes.
"""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

FONT = "Arial"
NAVY = "16325C"
HEAD = PatternFill("solid", fgColor=NAVY)
BAND = PatternFill("solid", fgColor="EEF2F7")
KEY = PatternFill("solid", fgColor="FFF4CC")
MONEY = '#,##0.00;(#,##0.00);"-"'
MONEY_ZERO = '#,##0.00;(#,##0.00);0.00'
THIN = Side(style="thin", color="C9D3E0")

IMPACT_MEANING = {
    "BREAK: payment confirmed, wallet not credited": "Owed to the player",
    "BREAK: unrecognised settlement (no internal record)": "Settled money not yet identified as ours",
    "BREAK: net amount is not gross minus fee": "Short-paid by the gateway",
    "BREAK: settled fee differs from contracted fee": "Fee overcharged by the gateway",
    "BREAK: settled gross amount differs from internal amount": "Internal minus gateway (negative: player under-credited)",
    "BREAK: duplicate gateway settlement for one reference (double-credit risk)": "0 per row; exposure in the bridge",
    "BREAK: duplicate internal SUCCESS deposit for one settlement (double-credit risk)": "0 per row; exposure in the bridge",
    "REVERSAL: gateway reversed/charged back after settlement": "Reversed amount to claw back",
    "BREAK: deposit SUCCESS, no gateway settlement found": "Credited to the player, no money received yet",
    "TIMING: settlement expected in next period (created near cut-off)": "Awaiting settlement next period",
    "TIMING: prior-period deposit settled at start of period": "Awaiting match to last period's deposit",
    "NOT A PROBLEM: rounding difference <= 1 cent": "Internal minus gateway",
    "NOT A PROBLEM: reference formatting differs (matched after normalising)": "None",
}

ASSUMPTIONS = [
    ("Period", "2026-09-01 00:00:00 to 2026-09-07 23:59:59 UTC inclusive. All timestamps are UTC; "
               "both files are NAD only (checked)."),
    ("Totals compared", "Internal total = SUCCESS deposits only (only SUCCESS credits a wallet). "
                        "Gateway total = SETTLED rows only; REVERSED rows are listed as reversals."),
    ("Matching key", "Our gateway_ref = the gateway's merchant_ref after upper-casing and removing everything "
                     "but letters and digits. 6 references differ only in formatting and are listed as 'not a problem'."),
    ("Fee rule", "Contract fee = ROUND(2% x gross + 1.00, 2). Any fee or net difference of 1 cent or more is a break."),
    ("Rounding", "A gross difference of exactly 1 cent is treated as rounding (immaterial); 2 cents or more is a "
                 "break. This is a materiality choice for Finance to confirm."),
    ("Cut-off window", "A SUCCESS deposit created in the last 15 minutes with no settlement, and a settlement in "
                       "the first 15 minutes whose reference is not in this period's deposits, are timing "
                       "differences. They are provisional until matched in the adjacent weeks' files."),
    ("Settlement lag", "Matched settlements arrive {median} after the deposit (median), {max} at most. The nearest "
                       "unexplained items to either boundary are 2 hours or more away (the 5 unsettled deposits "
                       "15+ hours before period end; the earliest unrecognised settlement 2 hours after the start), "
                       "so the classification does not depend on the exact cut-off window."),
    ("Reversals", "A REVERSED row means the gateway pulled the money back after settling. Treated as a genuine "
                  "break until the matching wallet debit is confirmed; wallet postings are not in these files."),
    ("Duplicates", "Two SUCCESS deposits for one reference created seconds apart = one payment recorded twice; "
                   "the later is the duplicate. Two settlement rows with the same gateway_txn_id = one "
                   "transaction reported twice; the later is the duplicate."),
    ("Bridge basis", "The bridge is on gross amounts, which is what deposits record. Fee and net differences do "
                     "not move gross, so they are actions, not bridge lines."),
    ("Not listed", "FAILED deposits with no settlement agree on both sides (nothing paid, nothing settled), so "
                   "they are not differences."),
]


def _header(ws, row, values, widths=None):
    for i, v in enumerate(values, 1):
        c = ws.cell(row=row, column=i, value=v)
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.fill = HEAD
        c.alignment = Alignment(vertical="center", wrap_text=True)
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w


def _data_sheet(wb, title, df, widths, money_cols=()):
    ws = wb.create_sheet(title)
    _header(ws, 1, list(df.columns), widths)
    for r, rec in enumerate(df.itertuples(index=False), 2):
        for c, v in enumerate(rec, 1):
            cell = ws.cell(row=r, column=c, value=None if (isinstance(v, float) and v != v) else v)
            cell.font = Font(name=FONT)
            if df.columns[c - 1] in money_cols:
                cell.number_format = MONEY
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    return ws


def write_workbook(path, exc, dep, gw, bridge_lines, categories, classes, lags):
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"

    # --- data sheets ------------------------------------------------------------------------
    money_exc = ("financial_impact", "bridge_amount", "dep_amount", "gross_amount", "fee", "expected_fee", "net_amount")
    widths = {"exception_id": 11, "classification": 16, "category": 40, "explanation": 70, "next_step": 45,
              "owner": 20, "bridge_line": 36}
    ws_exc = _data_sheet(wb, "Exceptions", exc, [widths.get(c, 14) for c in exc.columns], money_exc)
    for row in ws_exc.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=exc.columns[cell.column - 1] in widths)
    _data_sheet(wb, "Internal deposits", dep, [12, 10, 20, 12, 9, 13, 14, 10], ("amount",))
    _data_sheet(wb, "Gateway settlement", gw, [14, 14, 20, 13, 9, 12, 9, 11], ("gross_amount", "fee", "net_amount"))

    col = {name: get_column_letter(i) for i, name in enumerate(exc.columns, 1)}
    dcol = {name: get_column_letter(i) for i, name in enumerate(dep.columns, 1)}
    gcol = {name: get_column_letter(i) for i, name in enumerate(gw.columns, 1)}
    n = len(exc) + 1
    rng = lambda c: f"Exceptions!${col[c]}$2:${col[c]}${n}"

    # --- Summary ------------------------------------------------------------------------------
    s = summary
    for letter, w in zip("ABCDEF", (62, 16, 12, 16, 44, 24)):
        s.column_dimensions[letter].width = w
    s["A1"] = "Payment gateway reconciliation, 1-7 September 2026 (UTC, NAD)"
    s["A1"].font = Font(name=FONT, bold=True, size=14, color=NAVY)
    s["A2"] = ("Every figure below is a formula on the other sheets. Change a row in Exceptions or the raw "
               "files and the bridge recalculates; the residual must stay 0.00.")
    s["A2"].font = Font(name=FONT, italic=True, color="52514E")

    _header(s, 4, ["Bridge on gross amounts: internal total to gateway total", "NAD"])
    r = 5
    s.cell(row=r, column=1, value="Internal SUCCESS deposits, total")
    s.cell(row=r, column=2, value=f"=SUMIFS('Internal deposits'!${dcol['amount']}:${dcol['amount']},"
                                  f"'Internal deposits'!${dcol['status']}:${dcol['status']},\"SUCCESS\")")
    start = r
    for line in bridge_lines:
        r += 1
        s.cell(row=r, column=1, value=line)
        s.cell(row=r, column=2, value=f"=SUMIFS({rng('bridge_amount')},{rng('bridge_line')},A{r})")
    r += 1
    s.cell(row=r, column=1, value="Gateway total implied by the bridge")
    s.cell(row=r, column=2, value=f"=SUM(B{start}:B{r - 1})")
    implied = r
    r += 1
    s.cell(row=r, column=1, value="Gateway SETTLED total, from the settlement file")
    s.cell(row=r, column=2, value=f"=SUMIFS('Gateway settlement'!${gcol['gross_amount']}:${gcol['gross_amount']},"
                                  f"'Gateway settlement'!${gcol['status']}:${gcol['status']},\"SETTLED\")")
    actual = r
    r += 1
    s.cell(row=r, column=1, value="Residual (unexplained): must be 0.00")
    s.cell(row=r, column=2, value=f"=ROUND(B{actual}-B{implied},2)")
    residual = r
    for rr in range(start, residual + 1):
        for cc in (1, 2):
            cell = s.cell(row=rr, column=cc)
            cell.font = Font(name=FONT, bold=rr in (start, implied, actual, residual),
                             color="008000" if cc == 2 else "000000")
            cell.border = Border(bottom=THIN)
        s.cell(row=rr, column=2).number_format = MONEY_ZERO if rr == residual else MONEY
    s.cell(row=residual, column=1).fill = KEY
    s.cell(row=residual, column=2).fill = KEY

    r = residual + 2
    _header(s, r, ["Category", "Classification", "Rows", "Impact (NAD)", "What the impact means", "Owner"])
    first = r + 1
    for cat, (cls, _line, _step, owner) in categories.items():
        r += 1
        s.cell(row=r, column=1, value=cat)
        s.cell(row=r, column=2, value=cls)
        s.cell(row=r, column=3, value=f"=COUNTIFS({rng('category')},A{r})")
        s.cell(row=r, column=4, value=f"=SUMIFS({rng('financial_impact')},{rng('category')},A{r})")
        s.cell(row=r, column=5, value=IMPACT_MEANING[cat])
        s.cell(row=r, column=6, value=owner)
        for cc in range(1, 7):
            cell = s.cell(row=r, column=cc)
            cell.font = Font(name=FONT, color="008000" if cc in (3, 4) else "000000")
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if (r - first) % 2:
                cell.fill = BAND
        s.cell(row=r, column=4).number_format = MONEY
    last = r
    r += 1
    s.cell(row=r, column=1, value="Total listed differences")
    s.cell(row=r, column=3, value=f"=SUM(C{first}:C{last})")
    s.cell(row=r, column=1).font = s.cell(row=r, column=3).font = Font(name=FONT, bold=True)

    r += 2
    _header(s, r, ["Classification", "Rows"])
    for cls in classes:
        r += 1
        s.cell(row=r, column=1, value=cls).font = Font(name=FONT)
        s.cell(row=r, column=2, value=f"=COUNTIFS({rng('classification')},A{r})").font = Font(name=FONT, color="008000")
    s.freeze_panes = "A4"

    # --- Assumptions --------------------------------------------------------------------------
    a = wb.create_sheet("Assumptions")
    _header(a, 1, ["Topic", "Assumption"], [22, 120])
    for i, (topic, text) in enumerate(ASSUMPTIONS, 2):
        text = text.format(**lags)
        a.cell(row=i, column=1, value=topic).font = Font(name=FONT, bold=True)
        cell = a.cell(row=i, column=2, value=text)
        cell.font = Font(name=FONT)
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    wb.move_sheet("Assumptions", offset=-3)       # Summary, Assumptions, Exceptions, raw files
    wb.save(path)
